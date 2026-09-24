"""Versioned application drafts built from trusted profile evidence only.

Job text can select/reorder evidence; it can never supply candidate claims.
No send, submit, shell-from-model, or modification of source CVs is supported.
"""
from __future__ import annotations

import importlib.util
import re
import shutil
import subprocess
import tempfile
import copy
import json
from pathlib import Path

from careerops.store import digest, now
from careerops.draft_selection import choose_evidence

PROMPT_VERSION = "evidence-draft-v6-private-profile"
ALLOWED = {"DIRECT", "TRANSFERABLE", "INDIRECT", "CURRENT_PROJECT", "USER_PROVIDED", "VERIFIED", "PASS"}
SENSITIVE = re.compile(r"\b(?:ADHD|autism|disability|health condition|date of birth|ethnicity|religion|sexual orientation)\b", re.I)


def approved_evidence(profile):
    return [r for r in profile.get("evidence", []) if isinstance(r, dict)
            and str(r.get("status", "")).upper() in ALLOWED and r.get("id") and r.get("text")
            and r.get("automatic_use_allowed") is not False
            and not r.get("sensitive") and r.get("include_in_model") is not False
            and not SENSITIVE.search(r["text"])]


def relevant(records, job, maximum):
    return choose_evidence(records, job, maximum)


def record_refs(record):
    return [str(i) if str(i).startswith("bank:") else "bank:" + str(i) for i in (record.get("direct_evidence_ids") or record.get("evidence_ids") or [])]


def evidence_paragraphs(records, group_size=3):
    return [" ".join(r["text"].strip() for r in records[i:i + group_size])
            for i in range(0, len(records), group_size)]


def education_section(profile):
    """Print only approved source wording for explicitly completed qualifications."""
    evidence = {r["id"]: r["text"].strip() for r in approved_evidence(profile)}
    ids = []
    for qualification in profile.get("qualifications", []):
        if not isinstance(qualification, dict) or qualification.get("completed") is not True:
            continue
        refs = qualification.get("evidence_ids") or [qualification.get("evidence_id")]
        for ref in refs:
            key = ref if ref in evidence else "bank:" + str(ref)
            if key in evidence and key not in ids:
                ids.append(key)
    return {"heading": "Education", "paragraphs": [evidence[i] for i in ids], "evidence_ids": ids}


def _cv_text(material):
    lines = [material["name"], material["location"], material["contact_text"]]
    for section in material["sections"]:
        lines += ["", section["heading"]]
        if section.get("subheading"):
            lines.append(section["subheading"])
        lines.extend(section["paragraphs"])
    return "\n".join(lines)


def make_draft(job, profile):
    evidence = approved_evidence(profile)
    index = {r["id"]: r for r in evidence}
    warnings = list(profile.get("unresolved", []))
    if not profile.get("employment"):
        warnings.append("Employment history has not been imported. Add verified history before treating this CV as ready to apply.")
    if not profile.get("contact"):
        warnings.append("Contact details need to be added to the private profile before application.")
    rejected = [r for r in profile.get("evidence", []) if isinstance(r, dict) and r not in evidence]
    if rejected:
        warnings.append(f"{len(rejected)} unresolved, unsupported or sensitive evidence records were excluded from this draft.")
    gaps = job.get("evaluation", {}).get("gaps", [])
    blockers = job.get("evaluation", {}).get("blockers", [])
    if blockers:
        warnings.append("This job has confirmed blockers: " + "; ".join(map(str, blockers)))
    selected = relevant(evidence, job, 5)
    bullets = [{"text": r["text"], "evidence_ids": [r["id"]]} for r in selected]
    sections = []
    if selected:
        sections.append({"heading": "Relevant strengths", "paragraphs": [r["text"] for r in selected[:3]], "evidence_ids": [r["id"] for r in selected[:3]]})
    skills = []
    for skill in profile.get("skills", []):
        name = skill if isinstance(skill, str) else skill.get("name", "")
        matching = [r for r in evidence if name and name.lower() in r["text"].lower()]
        if matching:
            skills.append((name, matching[0]["id"]))
    if skills:
        skills.sort(key=lambda s: s[0].lower() in job.get("description", "").lower(), reverse=True)
        sections.append({"heading": "Technical skills", "paragraphs": [" · ".join(s[0] for s in skills[:20])], "evidence_ids": [s[1] for s in skills[:20]]})
    # Preserve strict reverse chronology for employment; unknown dates go last.
    def year(record):
        raw = str(record.get("end", ""))
        if raw.lower() in {"present", "current", "ongoing", "now"}:
            return (9999, 12, 31)
        match = re.search(r"((?:19|20)\d\d)(?:-(\d{2})(?:-(\d{2}))?)?", raw)
        if not match:
            return (0, 0, 0)
        month = int(match.group(2) or 0)
        months = ["january", "february", "march", "april", "may", "june", "july", "august", "september", "october", "november", "december"]
        for number, name in enumerate(months, 1):
            if re.search(r"\b" + name[:3] + r"(?:" + name[3:] + r")?\b", raw, re.I):
                month = number
        return (int(match.group(1)), month, int(match.group(3) or 0))
    for record in sorted(profile.get("employment", []), key=year, reverse=True):
        refs = record_refs(record)
        records = relevant([index[r] for r in refs if r in index], job, 9)
        title = f"{record.get('title', 'Role')} | {record.get('employer', 'Employer')}"
        dates = " – ".join(str(record.get(k) or "Date to verify") for k in ("start", "end"))
        sections.append({"heading": title, "subheading": dates, "paragraphs": evidence_paragraphs(records), "evidence_ids": [r["id"] for r in records], "profile_ref": record.get("record_id", "employment")})
    # Project and research prose is copied from evidence, never commercialized.
    for category, maximum in (("projects", 3), ("research", 1)):
        for record in profile.get(category, [])[:maximum]:
            records = relevant([index[r] for r in record_refs(record) if r in index], job, 9)
            if not records:
                continue
            sections.append({"heading": record.get("name", category.title()), "category": category, "subheading": ("Personal / academic project" if category == "projects" else "Academic research") + " · " + str(record.get("period", "Dates to verify")), "paragraphs": evidence_paragraphs(records), "evidence_ids": [r["id"] for r in records], "profile_ref": record.get("record_id")})
    sections.append(education_section(profile))
    name = profile.get("name", "")
    contact = profile.get("contact", {})
    contact_text = " | ".join(str(contact.get(k, "")) for k in ("email", "phone", "linkedin_url") if contact.get(k))
    cv_text = _cv_text({"name": name, "location": str(profile.get("location", "")),
                        "contact_text": contact_text, "sections": sections})
    questions = job.get("application_questions", [])
    answers = [{"question": q if isinstance(q, str) else q.get("question", ""), "answer": "", "status": "needs_answer"} for q in questions]
    if answers:
        warnings.append("Known application questions require answers; blank answers remain visible in the preview.")
    prep = ["Prepare a specific example for each evidence bullet; explain your own contribution and its limits.",
            "Be ready to explain how your experience relates to this role's responsibilities.",
            "Confirm advertised requirements and practical eligibility before submitting."]
    prep += ["Resolve before applying: " + str(g) for g in gaps[:3]]
    pack = [f"Application draft for {job['title']} at {job['company']}", f"Source: {job.get('url') or 'pasted description'}", "", "REVIEW BEFORE APPLYING", *warnings, "", "CV DRAFT", cv_text, "", "EVIDENCE FOR THIS ROLE"]
    pack += [f"- {b['text']} [{', '.join(b['evidence_ids'])}]" for b in bullets]
    pack += ["", "INTERVIEW PREPARATION", *["- " + t for t in prep]]
    if answers:
        pack += ["", "APPLICATION QUESTIONS", *[a["question"] + "\nAnswer needed" for a in answers]]
    return {"created_at": now(), "status": "needs_review", "text": "\n".join(pack), "cv_text": cv_text, "sections": sections,
            "name": name, "location": str(profile.get("location", "")), "contact_text": contact_text,
            "evidence_bullets": bullets, "interview_notes": prep, "answers": answers, "warnings": warnings,
            "provenance": {r["id"]: {"source": r.get("source"), "text": r["text"]} for r in evidence if any(r["id"] in s.get("evidence_ids", []) for s in sections)},
            "provider": "deterministic_evidence_selection", "prompt_version": PROMPT_VERSION,
            "formats": ["txt", "docx"] + (["pdf"] if soffice_path() else [])}


def family_base(store, profile, family):
    """Reuse a verified profile base for each family; source facts stay verbatim."""
    key = digest([profile, family, PROMPT_VERSION])
    with store.lock, store.connect() as db:
        row = db.execute("SELECT data FROM cv_bases WHERE cache_key=?", (key,)).fetchone()
        if row:
            return json.loads(row[0])
        base = copy.deepcopy(profile)
        target = {"title": family.replace("_", " "), "description": family.replace("_", " ")}
        evidence = approved_evidence(profile)
        selected = choose_evidence(evidence, target, min(60, len(evidence)))
        selected_ids = {r["id"] for r in selected}
        base["evidence"] = selected + [r for r in profile.get("evidence", []) if r.get("id") not in selected_ids]
        base["cv_base"] = {"key": key, "family": family, "profile_version": profile.get("version"), "policy_version": PROMPT_VERSION,
                           "verified_evidence_ids": [r["id"] for r in evidence], "created_at": now()}
        db.execute("INSERT INTO cv_bases VALUES (?,?)", (key, json.dumps(base)))
    return base


def validate_draft(material, profile):
    if material.get("document_schema"):
        from .cv_document import validate_document
        return validate_document(material, profile)
    if material.get("provider") != "deterministic_evidence_selection":
        raise ValueError("Only evidence-grounded generated documents pass automatic validation.")
    approved = {r["id"]: r["text"] for r in approved_evidence(profile)}
    education = [s for s in material.get("sections", []) if s.get("heading") == "Education"]
    if education != [education_section(profile)]:
        raise ValueError("Protected qualifications changed or lack approved evidence.")
    for section in material.get("sections", []):
        ids = section.get("evidence_ids", [])
        for evidence_id in ids:
            if evidence_id not in approved:
                raise ValueError("Draft references unsupported candidate evidence.")
        if section.get("heading") == "Technical skills":
            names = {s if isinstance(s, str) else s.get("name", "") for s in profile.get("skills", [])}
            for paragraph in section.get("paragraphs", []):
                for name in paragraph.split(" · "):
                    if name not in names or not any(name.lower() in approved[eid].lower() for eid in ids):
                        raise ValueError("An exported skill lacks profile evidence.")
            continue
        for paragraph in section.get("paragraphs", []):
            remaining = paragraph
            for evidence_id in sorted(ids, key=lambda i: len(approved[i]), reverse=True):
                remaining = remaining.replace(approved[evidence_id].strip(), "")
            if remaining.strip():
                raise ValueError("A candidate claim does not trace to the selected evidence.")
    if material.get("cv_text") != _cv_text(material):
        raise ValueError("Document text differs from its validated sections.")
    return {"status": "passed", "checked_at": now(), "policy_version": PROMPT_VERSION, "evidence_claims": len(material.get("provenance", {}))}


def prepare(store, job_id, family=None, *, refresh=True):
    job = store.get_job(job_id)
    profile = store.profile()
    if family is None:
        from careerops.inventory import classify_job
        family = classify_job(job, store.settings())["role_family"]
    key = digest([job.get("content_fingerprint"), job["id"], profile, PROMPT_VERSION, family,
                  job.get('evaluation', {}).get('policy_version')])
    with store.connect() as db:
        cached = db.execute("SELECT id FROM materials WHERE cache_key=?", (key,)).fetchone()
    if cached:
        return store.material(cached[0])
    base = family_base(store, profile, family)
    draft = make_draft(job, base)
    draft["cv_base"] = base["cv_base"]
    draft["validation"] = validate_draft(draft, profile)
    draft["version"] = len(job.get("materials", [])) + 1
    result = store.save_material(job_id, key, draft)
    store.action(job_id, "materials_ready", refresh=refresh)
    return result


def edit_draft(store, material_id, text):
    original = store.material(material_id)
    if not isinstance(text, str) or not text.strip() or len(text) > 100000:
        raise ValueError("Draft text must contain 1–100,000 characters.")
    revised = dict(original, text=text, status="needs_review", formats=["txt"], version=original["version"] + 1,
                   created_at=now(), warnings=original.get("warnings", []) + ["Manual text edits have not been evidence validated. Only the editable draft text is exported for this revision."], parent_id=material_id)
    revised.pop("id", None)
    result = store.save_material(original["job_id"], digest([material_id, text]), revised)
    store.action(original["job_id"], "materials_ready")
    return result


def soffice_path():
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.is_file() else None


def render_docx(material, target):
    from docx import Document
    from docx.shared import Inches, Mm, Pt, RGBColor
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn

    doc = Document()
    title_properties = doc.styles["Title"].element.pPr
    if title_properties is not None:
        for border in list(title_properties.findall(qn("w:pBdr"))):
            title_properties.remove(border)
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(16), Mm(15)
    section.left_margin, section.right_margin = Mm(18), Mm(18)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.04
    for name in ("Heading 1", "Heading 2"):
        doc.styles[name].font.name = "Calibri"
        doc.styles[name].font.size = Pt(11)
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
        doc.styles[name].paragraph_format.space_before = Pt(9)
        doc.styles[name].paragraph_format.space_after = Pt(3)
        doc.styles[name].paragraph_format.keep_with_next = True
    title = doc.add_paragraph(material["name"], "Title")
    title.paragraph_format.space_after = Pt(3)
    for run in title.runs:
        run.font.size, run.font.color.rgb = Pt(23), RGBColor(0, 0, 0)
    if material.get("positioning"):
        p = doc.add_paragraph(material["positioning"])
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(11)
    doc.add_paragraph(" | ".join(v for v in [material.get("location"), material.get("contact_text")] if v))
    for s in material["sections"]:
        doc.add_paragraph(s["heading"], "Heading 2")
        if s.get("subheading"):
            p = doc.add_paragraph(s["subheading"])
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(9)
        for i, text in enumerate(s["paragraphs"]):
            p = doc.add_paragraph(text)
            if s.get("category") and i < len(s["paragraphs"]) - 1:
                p.paragraph_format.keep_with_next = True
    doc.core_properties.title = "Curriculum vitae"
    doc.core_properties.author = material["name"]
    doc.core_properties.subject = "Application draft"
    doc.save(target)


def export_material(store, material_id, format):
    material = store.material(material_id)
    if format not in material["formats"]:
        raise ValueError("That export is unavailable for this draft revision.")
    folder = store.path.parent / "materials" / str(material_id)
    folder.mkdir(parents=True, exist_ok=True, mode=0o700)
    target = folder / ("application-draft." + format)
    if target.exists():
        if format != "docx":
            return target
        try:
            from docx import Document
            Document(target)
            return target
        except Exception:
            # An interrupted export must be recoverable on batch resume.
            target.unlink()
    if format == "txt":
        target.write_text(material["text"], encoding="utf-8")
    elif format == "docx":
        from docx import Document
        with tempfile.TemporaryDirectory(prefix=".export-", dir=folder) as tmp:
            staged = Path(tmp) / target.name
            render_docx(material, staged)
            Document(staged)
            staged.chmod(0o600)
            staged.replace(target)
    elif format == "pdf":
        docx = export_material(store, material_id, "docx")
        with tempfile.TemporaryDirectory(prefix="careerops-render-") as tmp:
            result = subprocess.run([soffice_path(), "-env:UserInstallation=" + (Path(tmp) / "profile").as_uri(), "--headless", "--convert-to", "pdf", "--outdir", tmp, str(docx)], capture_output=True, timeout=60)
            pdf = Path(tmp) / "application-draft.pdf"
            if result.returncode or not pdf.exists():
                raise ValueError("PDF rendering failed; the editable DOCX and text draft remain available.")
            shutil.copy2(pdf, target)
    target.chmod(0o600)
    return target
